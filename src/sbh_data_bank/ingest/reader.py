"""Reading a vendor file into typed rows.

Streaming throughout. Curve Monthly Services is 260 MB and 230,192 rows, Curve
Appointments 143 MB, and the Monthly Payments set is 1.29 million rows across
40 files - none of which the previous pipeline ever loaded, and all of which it
would have processed a row at a time through ``df.iterrows()``. Nothing here
holds a whole file in memory.

Three invariants, taken from the definition of done:

**One bad row never fails a file.** A row whose values will not parse goes to
quarantine with a reason code and its raw payload. The file keeps loading.

**Every row is accounted for.** ``rows_read == rows_ok + rows_quarantined``,
asserted at the end of every load. A file that cannot account for every row it
read is a failure, not a partial success. This is the check that would have
caught the 227,284 payment rows the previous currency parser silently nulled.

**Prohibited values never materialise.** A ``PROHIBITED`` column is located so
that drift and column-count checks still see it, and then never read out of the
row. Its value does not reach the typed row, the quarantine payload, or a log
line - including when some *other* column on the same row fails to parse.

How this fits
-------------
This is the lowest layer of ingestion. It knows how to open a CSV or a
worksheet, find the header, hand back a stream of raw rows, and turn one raw
row into a typed row (or a quarantined one) using the binding that
``registry.binding`` produced and the parsers in ``transforms``. It does not
choose a variant, does not count rows for a whole file, and does not know
about footers or databases - those belong to ``ingest.load``, which is the
only module in ``src`` that imports this one. ``scripts/bind_real_files.py``
also uses ``open_source`` directly to probe real extracts, and
``tests/unit/test_load.py`` exercises it end to end on small fixtures.

The data types defined here (``RawRow``, ``TypedRow``, ``QuarantinedRow``,
``CellError``, ``ReadCounters``) are what ``ingest.load`` collects into a
``FileLoad``, so reading this file first makes ``load.py`` easy to follow.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path
from typing import Any, BinaryIO

from sbh_data_bank.registry.binding import Binding
from sbh_data_bank.registry.headers import normalize_header
from sbh_data_bank.registry.spec import FieldType, OnError, PiiClass, ReadSpec
from sbh_data_bank.transforms.ndc import NdcDirectory, normalize_ndc
from sbh_data_bank.transforms.values import (
    TEXT_LENGTH,
    ParseResult,
    parse_bool,
    parse_currency,
    parse_date,
    parse_decimal,
    parse_int,
    parse_text,
)

__all__ = [
    "CellError",
    "HeaderNotFoundError",
    "QuarantinedRow",
    "RawRow",
    "ReadCounters",
    "TypedRow",
    "open_source",
    "resolve_sheet",
    "type_row",
]

#: File extensions that are opened as workbooks (via openpyxl) rather than as
#: delimited text. Anything else - .csv, .txt, a zip member stream - goes
#: through the csv module. The extension decides the *reader*, never the
#: *source*: source identity comes from the header fingerprint in the spec,
#: because the same report has arrived as CSV in one drop and XLSX in the next
#: (WeInfuse Patients, ABC).
_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}


@dataclass(frozen=True, slots=True)
class RawRow:
    """One row as it came out of the file, before typing.

    Produced by :func:`open_source`; consumed by :func:`type_row`. Values are
    kept exactly as the reader delivered them (strings from CSV, native cell
    values from openpyxl) so that the quarantine payload can show an operator
    what was actually in the file.
    """

    line_number: int
    """1-based position in the file, header included, so an operator can find
    the row in the original."""
    values: tuple[Any, ...]
    """Cells in file order. Indexed by the positions in the binding."""
    sheet: str | None = None
    """Worksheet name for workbook rows; ``None`` for CSV rows."""


@dataclass(frozen=True, slots=True)
class CellError:
    """Why one cell could not be used.

    One ``CellError`` per failing cell; a row collects all of them so an
    operator sees every problem at once rather than fixing them one load at a
    time. ``reason`` is a short code (``MISSING_REQUIRED_VALUE``,
    ``NDC_AMBIGUOUS``, ``DATE_FORMAT`` ...) that becomes the quarantine reason
    in ``stg.quarantine.reasons``.
    """

    column: str
    reason: str
    raw: str
    """The offending value. Never populated for a PROHIBITED column - see
    :func:`type_row`."""


@dataclass(frozen=True, slots=True)
class TypedRow:
    """A row that parsed cleanly.

    ``values`` is keyed by the spec's *target* names (``invoice_number``,
    ``ndc``), never by the vendor's header text, so everything downstream is
    independent of which variant the file was. ``ingest.load`` later adds a
    ``_row_hash`` entry to ``values`` for identity.
    """

    line_number: int
    values: dict[str, Any]
    sheet: str | None = None


@dataclass(frozen=True, slots=True)
class QuarantinedRow:
    """A row that did not parse, kept whole so it can be fixed and replayed.

    This is what lands in ``stg.quarantine``: the line number to find it, the
    errors that explain it, and the raw payload so it can be corrected and
    re-typed without going back to the vendor file.
    """

    line_number: int
    errors: tuple[CellError, ...]
    payload: dict[str, str]
    """Raw values by target name, for the columns the spec knows. Excludes every
    PROHIBITED column."""
    sheet: str | None = None

    @property
    def reasons(self) -> tuple[str, ...]:
        """The distinct reason codes on this row, sorted so they are stable in
        logs and in the ``reasons`` column of the quarantine table."""
        return tuple(sorted({e.reason for e in self.errors}))


def _open_text(source: Path | BinaryIO, read: ReadSpec) -> io.TextIOBase:
    """Open a CSV source as text using the spec's declared encoding.

    Two deliberate choices. ``errors="strict"`` so a file in the wrong
    encoding fails loudly instead of producing mojibake in identifiers that
    then silently fail to join. ``newline=""`` because the csv module handles
    line endings itself; without it an embedded newline inside a quoted cell
    (the Master COGS List has multi-line ship-to addresses in ``Location``)
    would be mangled. The default encoding on ``ReadSpec`` is ``utf-8-sig``,
    which strips the UTF-8 byte-order mark that Allscripts eRx ships, so the
    first header reads as ``IMREENC_CODE`` and not as BOM + ``IMREENC_CODE``.
    """
    if isinstance(source, Path):
        return source.open("r", encoding=read.encoding, errors="strict", newline="")
    # A stream (zip member, blob) is wrapped rather than re-opened. It is the
    # caller's handle, so it is also the caller's to close.
    return io.TextIOWrapper(source, encoding=read.encoding, errors="strict", newline="")


def _trim_trailing_empty(cells: list[Any]) -> list[str]:
    """Stringify a header row and drop empty cells from its right-hand end.

    Workbooks report a "used range" that can run past the real header: the
    Master COGS List's used range reaches column K while the header stops at
    G, so H-K come back as ``None`` (file-audit.md). Trimming them keeps the
    column count honest; otherwise every row would look like it had more
    fields than the header. Only *trailing* empties are removed - an empty
    header in the middle is a real problem and is left for the header matcher
    to flag.
    """
    out = [("" if c is None else str(c)) for c in cells]
    while out and not out[-1].strip():
        out.pop()
    return out


class HeaderNotFoundError(ValueError):
    """The declared header marker was not found near the top of the file.

    Loud on purpose. The alternative - falling back to "probably row 0" - is
    how a report title gets read as a header and every column binds to nothing.
    """


def _matches_marker(cells: list[Any], marker: str) -> bool:
    """Whether a candidate row contains the header-marker column.

    Compared on the normalized form (case, punctuation and whitespace folded)
    so ``Account #`` still matches if the vendor writes ``ACCOUNT #`` or
    ``Account#``. Any one cell matching is enough: the marker names one column
    that only the real header carries.
    """
    wanted = normalize_header(marker)
    return any(normalize_header(c) == wanted for c in cells)


def _csv_header_and_rows(
    source: Path | BinaryIO, read: ReadSpec
) -> tuple[list[str], Iterator[RawRow]]:
    """Find the header in a delimited file and return it with a row stream.

    Returns ``(header, rows)`` where ``header`` is the list of column names
    and ``rows`` is a lazy iterator of :class:`RawRow`. Raises
    :class:`HeaderNotFoundError` when a ``header_marker`` is declared and not
    found. The file is read exactly once: the header is consumed here, and the
    same ``csv.reader`` continues from the next line when the caller iterates
    ``rows``.
    """
    handle = _open_text(source, read)
    reader = csv.reader(handle, delimiter=read.delimiter)

    # Two ways to locate the header, chosen by the spec:
    #   header_marker   scan for the row containing a named column
    #   preamble_lines  skip a fixed number of lines, then take the next
    # The marker form exists because a fixed count is not always possible.
    if read.header_marker is not None:
        # Scan for the row that carries the marker. ABC delivers the same 26
        # columns with one title row above the header as a workbook and two
        # lines - a sep= declaration plus a title - as a CSV, so a fixed count
        # cannot serve both and the marker has to locate the header instead.
        header: list[str] = []
        skipped = 0
        for candidate in reader:
            skipped += 1
            if _matches_marker(candidate, read.header_marker):
                header = _trim_trailing_empty(candidate)
                break
            # Bounded scan. A marker that is genuinely absent must fail fast
            # rather than reading a 260 MB file to the end looking for it.
            if skipped >= read.max_preamble_scan:
                raise HeaderNotFoundError(
                    f"header marker {read.header_marker!r} not found in the first "
                    f"{read.max_preamble_scan} rows"
                )
        else:
            # for/else: the loop ran out of rows before `break` - the file is
            # shorter than max_preamble_scan and still has no header.
            raise HeaderNotFoundError(
                f"header marker {read.header_marker!r} not found - file ended first"
            )
        # `skipped` counts the header row itself, so it is the header's
        # 1-based line number.
        header_line = skipped
    else:
        # Fixed preamble. Athena reports have one title row ("REPORT NAME :"),
        # most other sources have none, so preamble_lines is 1 or 0.
        for _ in range(read.preamble_lines):
            handle.readline()
        header = _trim_trailing_empty(next(reader, []))
        header_line = read.preamble_lines + 1

    def rows() -> Iterator[RawRow]:
        # Lazy generator over the rest of the file. Line numbers continue from
        # the header so they match what an operator sees in a text editor.
        # The handle is closed when the generator is exhausted or abandoned,
        # but only if this module opened it (a Path); a caller-supplied stream
        # stays open for the caller.
        try:
            for offset, cells in enumerate(reader, start=1):
                yield RawRow(line_number=header_line + offset, values=tuple(cells))
        finally:
            if isinstance(source, Path):
                handle.close()

    return header, rows()


def _excel_header_and_rows(
    path: Path, read: ReadSpec, sheet: str
) -> tuple[list[str], Iterator[RawRow]]:
    """The workbook counterpart of :func:`_csv_header_and_rows`.

    Same contract: ``(header, rows)`` for one named worksheet, header located
    by marker or by fixed preamble, :class:`HeaderNotFoundError` when a marker
    is declared and absent. ``KeyError`` propagates from openpyxl if the sheet
    name does not exist; ``load_file`` treats that as "this variant's read
    settings do not fit" and tries the next variant.
    """
    # Imported here rather than at module level so that CSV-only use (and the
    # test suite's fast paths) never pay for loading openpyxl.
    import openpyxl

    # read_only streams rows instead of building the whole sheet in memory;
    # data_only returns cached formula results rather than formula text.
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    stream = worksheet.iter_rows(values_only=True)

    if read.header_marker is not None:
        # Identical scan to the CSV path. ABC's workbook form has one title
        # row above the header; the marker finds it without a fixed count.
        header = []
        skipped = 0
        for candidate in stream:
            skipped += 1
            cells = list(candidate)
            if _matches_marker(cells, read.header_marker):
                header = _trim_trailing_empty(cells)
                break
            if skipped >= read.max_preamble_scan:
                # Close before raising: read_only workbooks hold a file handle.
                workbook.close()
                raise HeaderNotFoundError(
                    f"header marker {read.header_marker!r} not found in the first "
                    f"{read.max_preamble_scan} rows of sheet {sheet!r}"
                )
        else:
            workbook.close()
            raise HeaderNotFoundError(
                f"header marker {read.header_marker!r} not found in sheet {sheet!r}"
            )
        header_line = skipped
    else:
        for _ in range(read.preamble_lines):
            next(stream, None)
        # `next(stream, ())` rather than `next(stream)`: an empty sheet yields
        # an empty header and lets the binder report "no columns" instead of a
        # bare StopIteration escaping from here.
        header = _trim_trailing_empty(list(next(stream, ())))
        header_line = read.preamble_lines + 1

    def rows() -> Iterator[RawRow]:
        # Each row records its sheet, because a workbook spec can read several
        # tabs and the natural key of the Master COGS List includes the tab.
        try:
            for offset, cells in enumerate(stream, start=1):
                yield RawRow(line_number=header_line + offset, values=tuple(cells), sheet=sheet)
        finally:
            workbook.close()

    return header, rows()


def resolve_sheet(path: Path, read: ReadSpec) -> list[str]:
    """Worksheet names this read spec selects, in workbook order.

    ``ReadSpec.selects_sheet`` applies the spec's ``sheet`` (exact,
    case-insensitive) or ``sheet_pattern`` (regex) rule; with neither set,
    every sheet is selected. Case-insensitive matters: the same report has
    arrived with tabs named ``Results``, ``sheet1`` and ``Sheet1``.
    ``ingest.load_workbook`` calls this once per variant to decide which tabs
    to load.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [n for n in workbook.sheetnames if read.selects_sheet(n)]
    finally:
        workbook.close()


def _is_blank(row: RawRow) -> bool:
    """True when every cell is ``None`` or whitespace.

    Used twice: to hold back trailing blank rows in :func:`open_source`, and
    to report an interior blank row as a single ``BLANK_ROW`` quarantine in
    :func:`type_row`.
    """
    return all(v is None or str(v).strip() == "" for v in row.values)


def open_source(
    source: Path | BinaryIO,
    read: ReadSpec,
    *,
    sheet: str | None = None,
) -> tuple[list[str], Iterator[RawRow]]:
    """Return the header row and a stream of data rows, in a single pass.

    One pass matters for a 260 MB CSV and for reading a member of a zip, where
    the handle cannot be rewound to read the header separately.

    Trailing blank rows are dropped when the spec asks - Cardinal detail exports
    end with one - but only *trailing* ones. A blank row in the middle of a file
    is a real anomaly and is passed through so it quarantines and somebody looks
    at it.

    Raises ``ValueError`` when a workbook path is given without a ``sheet``,
    and :class:`HeaderNotFoundError` from the underlying reader when a
    declared header marker is absent.
    """
    # Dispatch on extension. A stream (BinaryIO) is always treated as text,
    # because a workbook cannot be streamed without a path anyway.
    if isinstance(source, Path) and source.suffix.lower() in _EXCEL_SUFFIXES:
        if sheet is None:
            raise ValueError("reading a workbook requires a sheet name")
        header, rows = _excel_header_and_rows(source, read, sheet)
    else:
        header, rows = _csv_header_and_rows(source, read)

    if not read.skip_trailing_blank_rows:
        return header, rows

    def without_trailing_blanks() -> Iterator[RawRow]:
        # Blanks are held back rather than buffering the file: they are emitted
        # only once a non-blank row proves they were interior, and discarded at
        # end of file.
        #
        # Why this shape: we cannot know a blank row is trailing until the file
        # ends, and we cannot read the file twice. So blanks accumulate in
        # `pending`; the next real row flushes them (they were interior, so
        # they go through and quarantine as BLANK_ROW); end of file with
        # blanks still pending means they were trailing and they are dropped.
        # Memory cost is the length of the longest run of blank rows, not the
        # file.
        pending: list[RawRow] = []
        for row in rows:
            if _is_blank(row):
                pending.append(row)
                continue
            if pending:
                yield from pending
                pending.clear()
            yield row

    return header, without_trailing_blanks()


def _parse_cell(value: Any, column, ndc_directory: NdcDirectory | None) -> ParseResult[Any]:
    """Route one cell to the parser for its declared :class:`FieldType`.

    Returns a :class:`ParseResult` in every case: ``ok`` with a typed value
    (or ``None`` for an empty cell), or failed with a reason code. Never
    raises on bad data; the ``AssertionError`` at the end can only fire if a
    new ``FieldType`` is added to the enum without a parser here, which is a
    programming error rather than a data problem.
    """
    kind = column.type
    if kind is FieldType.TEXT:
        # The limit is the width of the typed column: the spec's max_length,
        # or TEXT_LENGTH when it declares none. Longer values fail the row
        # rather than truncate (truncating an identifier corrupts it) and
        # rather than the file (which is what the insert would do).
        return parse_text(value, max_length=column.max_length or TEXT_LENGTH)
    if kind is FieldType.INTEGER:
        return parse_int(value)
    if kind is FieldType.DECIMAL:
        return parse_decimal(value)
    if kind is FieldType.CURRENCY:
        # Handles "$3,767.71 " and "($40.00)". Decimal, never float.
        return parse_currency(value)
    if kind is FieldType.BOOLEAN:
        return parse_bool(value)
    if kind is FieldType.DATE:
        # The spec's explicit format list is mandatory; nothing is inferred.
        return parse_date(value, column.date_formats)
    if kind is FieldType.NDC:
        # NDC has its own result type because resolution can need a directory
        # lookup (unhyphenated 10-digit input). Translate it into the common
        # ParseResult shape so the caller does not care.
        resolution = normalize_ndc(value, ndc_directory)
        if resolution.ok:
            return ParseResult(resolution.raw, resolution.ndc11)
        if not resolution.raw:
            # Empty cell: not an error. ABC has 367 service-charge lines with
            # no NDC that are real invoice lines (data-findings #9).
            return ParseResult.empty("")
        return ParseResult.failed(resolution.raw, resolution.quarantine_reason or "NDC_INVALID")
    raise AssertionError(f"unhandled field type {kind}")


def type_row(
    row: RawRow,
    binding: Binding,
    *,
    ndc_directory: NdcDirectory | None = None,
    injected: dict[str, Any] | None = None,
) -> TypedRow | QuarantinedRow:
    """Parse one raw row against its binding.

    ``injected`` supplies values for reader-injected columns - the worksheet
    name today - and for the spec's constants.

    A PROHIBITED column is never read. Its position is known, so drift checks
    work, but the value is not fetched, not typed, not put in the payload and
    not mentioned in an error. That holds even when a different column on the
    same row fails and the whole row is quarantined, which is the case that
    would otherwise leak an SSN into a quarantine table nobody thought of as
    sensitive.

    Returns a :class:`TypedRow` when every cell parsed and every required
    value is present, otherwise a :class:`QuarantinedRow` carrying every
    error found. Never raises on data; a row is either typed or quarantined.
    """
    injected = injected or {}
    values: dict[str, Any] = {}  # target name -> typed value (the row if it passes)
    errors: list[CellError] = []  # everything wrong with the row, not just the first
    payload: dict[str, str] = {}  # target name -> raw text (for the quarantine record)

    if _is_blank(row):
        # A wholly empty row reaching this point is interior - trailing ones are
        # dropped upstream. Reporting it as one BLANK_ROW is far more useful to
        # an operator than a dozen separate missing-value errors, and it names
        # the actual anomaly: a gap, a stray newline, or a totals footer.
        return QuarantinedRow(
            line_number=row.line_number,
            errors=(CellError(column="*", reason="BLANK_ROW", raw=""),),
            payload={},
            sheet=row.sheet,
        )

    # Walk the spec's columns (not the file's), so the output always has the
    # same keys regardless of which variant or column order the file used.
    for column_binding in binding.columns:
        column = column_binding.column

        if column.pii is PiiClass.PROHIBITED:
            # Skipped before the value is even indexed out of row.values. This
            # is the whole PHI guarantee: SSN in Monthly Payments and Curve
            # Monthly Services never leaves the raw tuple.
            continue

        if column.is_injected:
            # A column with no vendor source: the worksheet name, supplied by
            # the loader. Lands as given, no parsing.
            values[column.target] = injected.get(column.target)
            continue

        if column_binding.match is None:
            # Declared in the spec but absent from this file. Optional columns
            # land null; a *required* absent column was already caught by the
            # binder as drift and the file never reached here.
            values[column.target] = None
            continue

        # Short rows (fewer cells than the header) read as None past the end
        # rather than raising IndexError - a truncated last line is a data
        # problem to quarantine, not a crash.
        position = column_binding.match.position
        raw = row.values[position] if position < len(row.values) else None
        payload[column.target] = "" if raw is None else str(raw)

        result = _parse_cell(raw, column, ndc_directory)
        if result.ok:
            # `required` means the column exists AND this row has a value
            # (CLAUDE.md). Without the second half an interior blank-ish row
            # lands as a record with null everything.
            if column.required and result.value is None:
                errors.append(
                    CellError(column=column.target, reason="MISSING_REQUIRED_VALUE", raw="")
                )
                continue
            values[column.target] = result.value
            continue

        # Parse failure. on_error decides: NULL swallows it (explicit opt-in
        # only - this is the behaviour that lost 17% of every payments file
        # under the old parser), anything else records the error. FAIL is
        # treated like QUARANTINE at the row level; the file-level abort it
        # implies is not wired up in this layer yet.
        if column.on_error is OnError.NULL:
            values[column.target] = None
            continue
        errors.append(
            CellError(column=column.target, reason=result.reason or "UNPARSEABLE", raw=result.raw)
        )

    if errors:
        # The payload holds every known column's raw text except prohibited
        # ones, so the row can be corrected and replayed from the quarantine
        # table alone.
        return QuarantinedRow(
            line_number=row.line_number,
            errors=tuple(errors),
            payload=payload,
            sheet=row.sheet,
        )

    # Spec constants are stamped last, after the row is known to be good. This
    # is how one Pharmaforce spec serves two sites: each file's spec instance
    # injects its own site constant.
    for name, constant in binding.spec.constants.items():
        values[name] = constant
    return TypedRow(line_number=row.line_number, values=values, sheet=row.sheet)


def sum_column(rows: Iterator[TypedRow], column: str) -> Decimal:
    """Total one numeric column. Used for control totals.

    Non-Decimal values (None, text) are skipped rather than coerced, so a
    control total only ever sums values that parsed as money or quantity.
    Not in ``__all__``; ``ingest.load`` keeps its own running totals inline
    so the file is streamed once, and this helper exists for scripts and
    tests that already hold a list of rows.
    """
    total = Decimal(0)
    for row in rows:
        value = row.values.get(column)
        if isinstance(value, Decimal):
            total += value
    return total


@dataclass(slots=True)
class ReadCounters:
    """Row accounting for one file.

    Every row read must end up in exactly one bucket. A load that cannot say
    where each row went is not a partial success, it is a failure.

    Filled in by ``ingest.load._load_stream`` and persisted onto
    ``meta.batch`` (``rows_read``, ``rows_quarantined``, ``control_totals``)
    by the commit engine once it exists.
    """

    rows_read: int = 0
    """Data rows handed to the typer. Excludes the header and any recognised
    totals footer, which is checked rather than loaded."""
    rows_ok: int = 0
    """Rows that typed cleanly and are eligible to commit."""
    rows_quarantined: int = 0
    """Rows that failed typing and went to quarantine with reasons."""
    control_totals: dict[str, Decimal] = field(default_factory=dict)
    """Sum of each ``control_total: true`` column over the ok rows. Compared
    against the vendor's footer when the file has one."""

    @property
    def balances(self) -> bool:
        """The invariant: read == ok + quarantined. False means a bug in the
        loader, never a data problem, and the load must not commit."""
        return self.rows_read == self.rows_ok + self.rows_quarantined

    def imbalance(self) -> str | None:
        """A human-readable statement of the imbalance, or ``None`` when the
        counters balance. Becomes the batch ``detail`` on an UNBALANCED load."""
        if self.balances:
            return None
        return (
            f"row accounting does not balance: read {self.rows_read}, "
            f"ok {self.rows_ok} + quarantined {self.rows_quarantined} = "
            f"{self.rows_ok + self.rows_quarantined}"
        )
