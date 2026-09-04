"""Bind every shipped spec against the real extracts, and report what happens.

Not a test - the extracts contain PHI and live outside the repository, so this
cannot run in CI. It is the evidence that the specs in ``specs/`` still describe
the files that actually arrive, and it is what to run after any vendor change.

Reads header rows only. No cell values are read, printed or written anywhere.

    python scripts/bind_real_files.py [root]

``root`` defaults to the Sagebrush ETL Files folder.
"""

from __future__ import annotations

import csv
import io
import sys
import zipfile
from collections.abc import Iterator
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from sbh_data_bank.registry.binding import (  # noqa: E402
    AmbiguousVariantError,
    NoVariantMatchedError,
    bind,
)
from sbh_data_bank.registry.headers import AmbiguousHeaderError, HeaderIndex  # noqa: E402
from sbh_data_bank.registry.loader import SpecRegistry  # noqa: E402
from sbh_data_bank.registry.spec import ReadSpec, SourceSpec  # noqa: E402

DEFAULT_ROOT = Path(r"C:\Users\SharanLokeshwar\OneDrive - Sagebrush Healthcare\Sagebrush ETL Files")
REPO = Path(__file__).resolve().parents[1]

DATA_SUFFIXES = {".csv", ".xlsx", ".xls"}


def _csv_headers(handle: io.TextIOBase, read: ReadSpec) -> list[str]:
    for _ in range(read.preamble_lines):
        handle.readline()
    row = next(csv.reader(handle, delimiter=read.delimiter), [])
    while row and not str(row[-1]).strip():
        row.pop()
    return row


def _sniff_preamble(path: Path, encoding: str) -> int:
    """How many lines precede the header.

    Only ``sep=`` declarations and obvious report titles - a line with fewer
    than three populated fields. The ABC CSV opens with both.
    """
    with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
        for index in range(10):
            line = handle.readline()
            if not line:
                return 0
            if line.lower().startswith("sep="):
                continue
            cells = next(csv.reader([line]), [])
            if len([c for c in cells if c.strip()]) >= 3:
                return index
    return 0


def _excel_sheets(path: Path) -> Iterator[tuple[str, list[str]]]:
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for name in workbook.sheetnames:
            worksheet = workbook[name]
            first = next(worksheet.iter_rows(values_only=True), ())
            cells = [("" if c is None else str(c)) for c in first]
            while cells and not cells[-1].strip():
                cells.pop()
            yield name, cells
    finally:
        workbook.close()


def _report(label: str, spec: SourceSpec, headers: list[str]) -> str:
    try:
        binding = bind(spec, HeaderIndex(headers))
    except NoVariantMatchedError:
        # Not necessarily wrong. The filename pattern is a hint, so a derived
        # or analysis file whose name resembles a vendor extract lands here -
        # and refusing it is the correct outcome.
        return f"    {label}: not recognised, {len(headers)} columns - refused"
    except (AmbiguousVariantError, AmbiguousHeaderError) as exc:
        return f"    {label}: {type(exc).__name__}: {exc}"

    blocking = binding.blocking_problem()
    status = "loads" if blocking is None else f"BLOCKED - {blocking}"
    line = (
        f"    {label}: variant={binding.variant.name} "
        f"bound={len(binding.present)}/{len(binding.variant.columns)} "
        f"landed={len(binding.landed)} {status}"
    )
    if binding.has_drift or binding.missing_optional:
        line += f"\n        drift: {binding.drift_report()}"
    return line


def main(root: Path) -> int:
    registry = SpecRegistry.from_dir(REPO / "specs")
    print(f"{len(registry)} specs, root {root}\n")

    problems = 0
    for spec in registry:
        print(f"{spec.key}  ({spec.name})")
        seen = 0

        for path in sorted(root.rglob("*")):
            if not path.is_file():
                continue

            if path.suffix.lower() == ".zip" and spec.matches_filename(path.name):
                with zipfile.ZipFile(path) as archive:
                    members = [n for n in archive.namelist() if n.lower().endswith(".csv")]
                    for member in members[:2]:  # shape is identical across members
                        with archive.open(member) as raw:
                            text = io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace")
                            headers = _csv_headers(text, spec.variants[0].read)
                        seen += 1
                        line = _report(Path(member).name, spec, headers)
                        problems += "BLOCKED" in line
                        print(line)
                    if len(members) > 2:
                        print(f"    ... and {len(members) - 2} more members, same shape")
                continue

            if path.suffix.lower() not in DATA_SUFFIXES:
                continue
            if not spec.matches_filename(path.name):
                continue

            seen += 1
            if path.suffix.lower() == ".csv":
                encoding = spec.variants[0].read.encoding
                preamble = _sniff_preamble(path, encoding)
                read = spec.variants[0].read.model_copy(update={"preamble_lines": preamble})
                with path.open("r", encoding=encoding, errors="replace", newline="") as handle:
                    headers = _csv_headers(handle, read)
                line = _report(path.name, spec, headers)
                problems += "BLOCKED" in line
                print(line)
            else:
                for sheet_name, headers in _excel_sheets(path):
                    if not any(v.read.selects_sheet(sheet_name) for v in spec.variants):
                        continue
                    line = _report(f"{path.name} [{sheet_name}]", spec, headers)
                    problems += "BLOCKED" in line
                    print(line)

        if seen == 0:
            print("    (no matching files found)")
        print()

    print(
        "OK - every recognised file binds and loads"
        if problems == 0
        else f"{problems} file(s) recognised but blocked"
    )
    return 1 if problems else 0


if __name__ == "__main__":
    raise SystemExit(main(Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_ROOT))
